#!/usr/bin/env python3
"""Create an editable ACS-based JCTC manuscript outline, with native equations.

uv run --python /usr/bin/python3 --with python-docx --with latex2mathml \
    --with mathml2omml tools/create_second_order_outline.py

The official DOTX is downloaded separately into docs/templates. Never run this
generator over a manuscript that contains an author's edits or comments.
"""

import argparse
import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from latex2mathml.converter import convert as latex_to_mathml
from mathml2omml import convert as mathml_to_omml

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs/templates/acs_template_msw2010_ol_pc.dotx"
GUIDELINES = "https://researcher-resources.acs.org/publish/author_guidelines?coden=jctcce"
TEMPLATE_URL = "https://pubsapp.acs.org/paragonplus/submission/templates/acs_template_msw2010_ol_pc.dotx"


def load_template():
    """Convert the main content type from template to document in memory."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(TEMPLATE) as src, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as dst:
        for entry in src.infolist():
            data = src.read(entry.filename)
            if entry.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml")
            dst.writestr(entry, data)
    buffer.seek(0)
    doc = Document(buffer)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"):
            doc.element.body.remove(child)
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/comments")):
            del doc.part.rels[rel_id]
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path,
                        default=ROOT/"docs/jctc_second_order_outline.docx")
    args = parser.parse_args()
    if args.output.exists():
        raise SystemExit("Refusing to overwrite a potentially annotated manuscript; choose --output.")
    doc = load_template()
    for name in ("Normal", "TA_Main_Text", "BD_Abstract", "BB_Author_Name",
                 "BC_Author_Address", "TC_Table_Body", "VA_Figure_Caption"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")
        style.paragraph_format.line_spacing = 1.3
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.widow_control = True
    for name, size in (("Title", 20), ("BA_Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
        if name.startswith("Heading"):
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), str(int(name[-1])-1))
            style.element.get_or_add_pPr().append(outline)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(.8)
    cols = section._sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.set(qn("w:num"), "1")
    for container in (section.header, section.footer):
        for child in list(container._element):
            container._element.remove(child)
    footer = section.footer.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    doc.core_properties.title = "Matrix-Free Second-Order Optimization of Nonorthogonal Valence-Bond Wave Functions"
    doc.core_properties.author = "XMVB-CPP manuscript working group"
    doc.core_properties.subject = "JCTC Article outline for internal review; not a submission-ready manuscript"
    doc.core_properties.comments = "Adapted from the official ACS Word template; checked 2026-10-07."

    def body(text, style="TA_Main_Text"):
        p = doc.add_paragraph(text, style)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def h(text, level=1):
        return doc.add_heading(text, level=level)

    def note(paragraph, text):
        # Genuine Word comments; the recipient can reply or resolve them.
        if paragraph.runs:
            doc.add_comment(paragraph.runs, text=text, author="Internal review", initials="IR")

    equation_count = 0

    def eq(latex):
        nonlocal equation_count
        equation_count += 1
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width, table.columns[1].width = Inches(5.9), Inches(.6)
        p = table.cell(0, 0).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        omml = mathml_to_omml(latex_to_mathml(latex))
        math = parse_xml('<root xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                         +omml+"</root>")[0]
        p._p.append(math)
        number = table.cell(0, 1).paragraphs[0]
        number.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        number.add_run("(")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "SEQ Equation \\* ARABIC")
        run, value = OxmlElement("w:r"), OxmlElement("w:t")
        value.text = str(equation_count)
        run.append(value)
        field.append(run)
        number._p.append(field)
        number.add_run(")")
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(7)

    def table(caption, header, rows):
        body(caption, "VA_Figure_Caption")
        tab = doc.add_table(rows=1, cols=len(header))
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
            edge = OxmlElement("w:"+side)
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "4")
            edge.set(qn("w:color"), "B0B0B0")
            borders.append(edge)
        tab._tbl.tblPr.append(borders)
        for cell, text in zip(tab.rows[0].cells, header):
            cell.text = text
        repeat = OxmlElement("w:tblHeader")
        tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell, text in zip(tab.add_row().cells, row):
                cell.text = text
        for row in tab.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.style = doc.styles["TC_Table_Body"]
                    p.paragraph_format.line_spacing = 1.1
                    for run in p.runs:
                        run.font.size = Pt(10)
        body(" ")

    title = body(doc.core_properties.title, "BA_Title")
    note(title, "暂定题目。先以非正交 VBSCF 的 matrix-free 二阶框架为主线；是否将 NEO 作为主方法、TNHVP 作为对照，需结合统一版本的测试结果决定。")
    body("[Author names and author order to be confirmed]", "BB_Author_Name")
    body("[Affiliations; corresponding author e-mail; ORCID identifiers]", "BC_Author_Address")
    body("JCTC Article — manuscript outline for internal review, 7 October 2026")

    h("ABSTRACT")
    abstract = body(
        "Orbital optimization in valence-bond self-consistent-field theory must accommodate "
        "nonorthogonal orbitals, prescribed atomic-orbital supports, and relaxation of the structure "
        "coefficients. We formulate a matrix-free second-order treatment in support-preserving "
        "nonredundant coordinates and distinguish the relaxed-orbital and coupled orbital–structure "
        "models. The numerical study will compare truncated Newton and norm-extended optimization "
        "with conventional and block-preconditioned L-BFGS, reporting convergence, operator work, "
        "wall time, and peak memory under matched conditions. [Replace this sentence with the principal "
        "quantitative findings and the supported conclusion after the benchmark campaign.]",
        "BD_Abstract")
    note(abstract, "这是四句摘要框架，不是已完成研究的结果摘要。不能把‘外层步数减少’直接写成‘总计算更快’，也不能预先承诺所有体系都二次收敛。")

    h("1. INTRODUCTION")
    h("1.1. The optimization problem", 2)
    body("Explain why VBSCF uses nonorthogonal localized orbitals and, in HAO calculations, fixed "
         "orbital supports. Separate the cost of optimizing orbitals from the cost of solving the "
         "structure-coefficient problem. Introduce OEO as full-AO support, not as an orthogonal-orbital model.")
    h("1.2. Relation to second-order MCSCF optimization", 2)
    body("Review restricted-step Newton, augmented-Hessian and norm-extended optimization, and "
         "matrix-free CASSCF implementations [1–3]. Identify what carries over to VBSCF—the local "
         "coupled quadratic model and iterative operator actions—and what does not carry over without "
         "derivation: orthogonal orbital rotations and their redundancy rules.")
    h("1.3. Scope of this work", 2)
    body("State the methodological contribution as a second-order optimization framework for the "
         "specified nonorthogonal VB variational space. Treat integral representation and string "
         "contraction as supporting implementations, not independent novelty claims in this paper. "
         "The principal numerical question is when additional curvature information improves "
         "convergence enough to repay its cost relative to block-LBFGS.")

    h("2. THEORY")
    h("2.1. Stationary structure problem and state averaging", 2)
    body("Define the generalized structure eigenproblem at fixed orbital coefficients x. Use a "
         "different symbol for the AO overlap and the structure overlap throughout the manuscript.")
    eq(r"\mathbf H_{\mathrm{VB}}(\mathbf x)\mathbf a_s=E_s\mathbf S_{\mathrm{VB}}(\mathbf x)\mathbf a_s,\quad\mathbf a_r^T\mathbf S_{\mathrm{VB}}\mathbf a_s=\delta_{rs}")
    eq(r"\mathcal E(\mathbf x)=\sum_{s\in\mathcal A}w_sE_s(\mathbf x),\quad w_s=1/m")
    body("For the single-state case, use the lowest stationary eigenvalue; for equal-weight state "
         "averaging, define the selected eigenspace and its separation from excluded roots. Derive "
         "normalization and selected-subspace gauge projections before introducing response equations.")
    h("2.2. Support-preserving nonredundant orbital coordinates", 2)
    eq(r"\mathbf c_p=\mathbf R_p\mathbf x_p,\quad \mathbf x(\mathbf p)=\mathcal R_k(\mathbf p),\quad\mathcal R_k(\mathbf 0)=\mathbf x_k")
    body("Define the support injection R_p, horizontal lift U_k, local update map, and the physical "
         "step metric M_o. Classify a direction as redundant only if it leaves the represented "
         "wave function unchanged within the selected structure space. In particular, active–active "
         "mixing is not automatically redundant for an arbitrary selected-VB expansion.")
    eq(r"\mathbf g_o=\nabla_{\mathbf p}(\mathcal E\circ\mathcal R_k)(\mathbf 0),\quad \mathbf H_o\mathbf v=D^2(\mathcal E\circ\mathcal R_k)(\mathbf 0)[\mathbf v]")
    p = body("Give the derivative of the actual update map, rather than transferring an orbital-rotation "
             "Hessian to a coefficient update. Check the consistency of lift, norm, and retraction.")
    note(p, "最终稿需从冻结的代码版本核对 retraction 和 pullback。历史理论文档中的 additive chart 不能不加检查地当作当前实现。")
    h("2.3. Orbital–structure coupling and the relaxed Hessian", 2)
    body("At a structure-stationary accepted point, write the constrained local Hessian in orbital "
         "and projected structure tangent variables. A is the fixed-structure orbital block, B is "
         "the mixed derivative, and C is the structure tangent block. All blocks include the chosen "
         "normalization and generalized-overlap constraints.")
    eq(r"\mathbf z=\begin{pmatrix}\mathbf p\\\mathbf q\end{pmatrix},\quad\mathbf g=\begin{pmatrix}\mathbf g_o\\\mathbf 0\end{pmatrix},\quad\mathbf K=\begin{pmatrix}\mathbf A&\mathbf B^T\\\mathbf B&\mathbf C\end{pmatrix}")
    eq(r"\mathbf H_{\mathrm{rel}}=\mathbf A-\mathbf B^T\mathbf C^\dagger\mathbf B,\quad\mathbf q_N=-\mathbf C^\dagger\mathbf B\mathbf p_N")
    body("Define C† only on the admissible response space. Explain the equivalence of the unshifted "
         "coupled Newton and relaxed-orbital Newton limits, and why shifted coupled and eliminated "
         "models need not be identical away from that limit.")
    h("2.4. Restricted steps and residual certification", 2)
    eq(r"m(\mathbf z)=\mathcal E_k+\mathbf g^T\mathbf z+\tfrac12\mathbf z^T\mathbf K\mathbf z,\quad\mathbf z^T\mathbf M\mathbf z\leq\Delta_k^2")
    eq(r"\mathbf r=(\mathbf K+\lambda\mathbf M)\mathbf z+\mathbf g,\quad\lambda\geq0,\quad\mathbf K+\lambda\mathbf M\succeq0")
    eq(r"\lambda(\mathbf z^T\mathbf M\mathbf z-\Delta_k^2)=0,\quad\|\mathbf r\|_{\mathbf M^{-1}}\leq\epsilon_{\mathrm{abs}}+\eta_k\|\mathbf g\|_{\mathbf M^{-1}}")
    body("Separate stationarity, step feasibility, complementarity, and curvature checks. State "
         "the regularity and inner-solve conditions required for superlinear or quadratic local "
         "convergence; an HVP implementation alone does not establish those properties.")
    h("2.5. TNHVP and norm-extended optimization", 2)
    body("Describe TNHVP as a truncated Newton realization of the relaxed model, and NEO as a "
         "norm-extended formulation of the local second-order model [1]. Define the gradient scaling "
         "and the step extracted from the generalized augmented eigenvector explicitly.")
    eq(r"\begin{pmatrix}0&\alpha\mathbf g^T\\\alpha\mathbf g&\mathbf K\end{pmatrix}\begin{pmatrix}\beta\\\mathbf y\end{pmatrix}=\mu\begin{pmatrix}1&0\\0&\mathbf M\end{pmatrix}\begin{pmatrix}\beta\\\mathbf y\end{pmatrix}")
    eq(r"\mathbf z=\mathbf y/(\alpha\beta),\quad\lambda=-\mu,\quad(\mathbf K+\lambda\mathbf M)\mathbf z=-\mathbf g")
    body("Discuss the nonzero homogeneous component, lowest-root selection, and the interior "
         "Newton limit. Describe how the projected solver reaches this model; do not equate a "
         "generic augmented-Hessian step with a certified norm-constrained solution.")
    h("2.6. Preconditioning and controlled structure response", 2)
    p = body("Distinguish the Hessian model from the preconditioner. Present the one-electron block "
             "preconditioner, retained secant information where applicable, and the role of coupled "
             "structure response. If the implementation omits or restricts response according to "
             "a work estimate, identify that calculation as an approximate model and report it separately.")
    note(p, "当前 NEO 后端仍有 structure-response admission 分支。正文不能把所有生产运行都描述为完整 coupled Newton；必须分别记录 admitted / core-only 路径。")

    h("3. IMPLEMENTATION")
    h("3.1. Matrix-free derivative actions", 2)
    body("Describe the accepted-point data required for gradient and HVP evaluation: normalization, "
         "inactive projection, integral transformation, structure action, and response. Identify reused "
         "quantities and directional quantities, without claiming that matrix-free means no large "
         "intermediates. Give memory costs in orbital, structure, and unique-string dimensions.")
    h("3.2. Trial steps and accepted-point updates", 2)
    eq(r"\rho_k=\frac{\mathcal E_k-\mathcal E(\mathcal R_k(\mathbf p))}{m(\mathbf0)-m(\mathbf z)}")
    body("Explain trial evaluation, acceptance, trust-radius updates, and rebuilding the local "
         "model. State which quantities are refreshed at a keyframe and which are reused during "
         "subspace expansion. Count rejected trials as work even when they do not advance the "
         "accepted-iteration counter.")
    body("Algorithm 1. Accepted-point second-order VBSCF optimization (outline)", "VA_Figure_Caption")
    for text in (
        "1. Read the same orbital guess and structure list used by all comparison methods.",
        "2. Solve the selected structure problem; evaluate energy and the physical orbital gradient.",
        "3. Build the support-preserving local coordinates, metric, derivative actions, and preconditioner.",
        "4. Expand the iterative second-order model; certify its residual and record response treatment.",
        "5. Retract the orbital step, solve the trial structure problem, and evaluate model agreement.",
        "6. Accept or reject the trial; refresh accepted-point data and test both stopping criteria."):
        body(text)
    h("3.3. Cost accounting", 2)
    body("Report orbital HVPs, orbital-to-structure and transpose-coupling actions, structure "
         "H/S actions, trial evaluations, and accepted steps separately. Include setup, final "
         "observables, and total process memory in end-to-end measurements; isolate optimizer "
         "time in an additional breakdown.")

    h("4. COMPUTATIONAL DETAILS")
    h("4.1. Methods and matched conditions", 2)
    table("Table 1. Comparison methods. Input selectors are implementation details, not method names.",
          ["Method", "Selector", "Role"], [
              ("XMVB-style L-BFGS", "ISCF=5", "Conventional quasi-Newton baseline"),
              ("TNHVP", "ISCF=7", "Matrix-free truncated Newton comparison"),
              ("Block-LBFGS", "ISCF=8", "Enhanced baseline with one-electron block initial inverse"),
              ("NEO", "ISCF=9", "Norm-extended second-order method")])
    body("Use identical initial orbitals, geometry, basis, active electrons/orbitals, spin, "
         "selected structure list, state weights, integral representation, and final convergence "
         "tests. Compare exact and RI only in separate matched groups. Dense and Davidson structure "
         "solvers must solve the same selected generalized eigenproblem to controlled accuracy.")
    p = body("Record the projected-gradient convention and the physical dual-metric norm. The "
             "working energy-change and projected-gradient targets are 10⁻⁷ hartree and 10⁻³, "
             "respectively; confirm the precise implemented norms and final settings before collection. "
             "Use an independently tighter calculation for endpoint comparisons.")
    note(p, "用户此前认可的阈值作为工作设置；不能把不同坐标中的无穷范数当作完全等价的物理梯度。正式稿应报告实际采用的定义。")
    h("4.2. Test-set design", 2)
    table("Table 2. Proposed test groups; molecular specifications and eligibility remain to be frozen.",
          ["Group", "Examples", "Purpose"], [
              ("Small derivative oracles", "F₂; small selected-structure cases", "Finite differences, dense Hessian, endpoint checks"),
              ("Localized HAO", "241 (benzene), 240, 7975, 7963, YAMSAI", "Support constraints and convergence variation"),
              ("Transition metals", "MnF₂ and FeCl₂", "Conditioning, spin, and OEO/HAO behavior"),
              ("Larger response spaces", "CERRAS and LOFLEA", "Response cost and memory"),
              ("Equal-weight state average", "Small multistate cases; selected pyrrole/OEO cases", "Root-space and response stability")])
    p = body("Do not include unresolved endpoint discrepancies as successful validation. A PySCF "
             "CASSCF comparison is a variational-space reference only when the OEO structure expansion, "
             "spin/symmetry sector, integrals, and core convention match; otherwise it is a diagnostic comparison.")
    note(p, "FeCl₂、部分 pyrrole/OEO 的历史问题要在版本冻结时重新确认。这里仅列测试候选，不把历史能量差当成已解决。")
    h("4.3. Hardware and reproducibility", 2)
    body("Collect publication timings on Hanhai25 compute nodes through Slurm. Report CPU model, "
         "allocated cores, affinity, OpenMP/BLAS threads, compiler flags, dependency versions, and "
         "commit identifier. Record peak RSS at the job/process level. Preserve nonconverged runs "
         "and resource failures in the dataset rather than dropping them.")

    h("5. RESULTS AND DISCUSSION")
    h("5.1. Derivative and response validation", 2)
    body("Table 3 will report gradient directional differences, HVP directional differences, "
         "bilinear symmetry defects, and coupled-versus-Schur agreement on small cases. Include "
         "nonorthogonal well-conditioned and near-singular orbital points, HAO and OEO, and at least "
         "one equal-weight state average. [Insert verified numerical errors.]")
    h("5.2. Local convergence", 2)
    body("Figure 2 will show energy error and physical gradient norm for identical controlled "
         "perturbations around a common stationary point. Establish the local convergence regime "
         "separately from the number of iterations required to enter it. [Insert trajectories and "
         "the residual forcing sequence.]")
    h("5.3. End-to-end comparison", 2)
    body("Figures 3 and 4 will plot accepted steps and total wall time against a common system "
         "ordering, with failures visibly marked. Table 4 will provide final energies, energy "
         "changes, physical gradients, action counts, wall time, and RSS. Discuss cases where "
         "fewer steps do not yield lower wall time. [Insert matched results; do not mix revisions.]")
    h("5.4. Where the second-order cost is spent", 2)
    body("Figure 5 will break down setup, core HVP, coupling/structure response, trial evaluation, "
         "and final output. Analyze larger active spaces using measured work and memory rather than "
         "accepted-step count alone. Compare full and restricted response with the model treatment "
         "stated for every run. [Insert stage timings and response statistics.]")
    h("5.5. Ablations and limitations", 2)
    body("Separate model changes from preconditioner changes. Prioritize matched full/restricted "
         "response and one-electron block/Hessian-diagonal comparisons; identity preconditioning "
         "need not be a production benchmark. Discuss small structure gaps, incomplete selected-VB "
         "spaces, root changes, and redundant-coordinate conditioning. [Insert only ablations that "
         "answer a specific algorithmic question.]")

    h("6. CONCLUSIONS")
    body("Summarize the validated second-order formulation, the measured range of efficiency, and "
         "the cost of structure relaxation. State which method is recommended for which workload "
         "only if the data support that distinction. [Write the quantitative conclusion after "
         "Sections 5.1–5.5 are complete.]")
    h("ASSOCIATED CONTENT")
    h("Supporting Information", 2)
    body("Planned contents: coordinate and response derivations; update-map derivatives; "
         "finite-difference protocols; complete molecular specifications and initial orbitals; "
         "raw convergence trajectories; nonconvergence records; timing scripts and environment "
         "details. Submit the SI separately from the main manuscript.", "TE_Supporting_Information")
    h("AUTHOR INFORMATION")
    body("Corresponding Author: [Name and e-mail]. Author Contributions: [To be agreed by all authors]. "
         "Notes: [Competing-interest statement to be confirmed].")
    h("ACKNOWLEDGMENTS")
    body("[Funding, computing resources, and individual acknowledgments to be confirmed. "
         "Add any required AI-assistance disclosure after review of the applicable ACS policy; "
         "do not list an AI tool as an author.]", "TD_Acknowledgments")
    h("DATA AVAILABILITY STATEMENT")
    body("[Supply a persistent public repository link or DOI for the input files, initial orbitals, "
         "raw outputs, benchmark metadata, and analysis scripts. Identify the code release and "
         "commit used for every figure. Explain any justified restrictions.]")
    h("REFERENCES — STARTER LIST")
    for text in (
        "(1) Jensen, H.-J. Aa.; Jørgensen, P. A Direct Approach to Second-Order MCSCF Calculations "
        "Using a Norm Extended Optimization Scheme. J. Chem. Phys. 1984, 80, 1204–1214. "
        "https://doi.org/10.1063/1.446797.",
        "(2) Werner, H.-J.; Knowles, P. J. A Second Order Multiconfiguration SCF Procedure with "
        "Optimum Convergence. J. Chem. Phys. 1985, 82, 5053–5063. https://doi.org/10.1063/1.448627.",
        "(3) Sun, Q.; Yang, J.; Chan, G. K.-L. A General Second Order Complete Active Space "
        "Self-Consistent-Field Solver for Large-Scale Systems. Chem. Phys. Lett. 2017, 683, "
        "291–299. https://doi.org/10.1016/j.cplett.2017.03.004."):
        body(text)
    body("[Complete the VBSCF/XMVB, inexact-Newton, trust-region, L-BFGS, and density-fitting "
         "literature before submission. This starter list is not an exhaustive priority review.]")

    doc.add_page_break()
    h("INTERNAL DRAFT NOTES — REMOVE BEFORE SUBMISSION")
    body("This file is an outline, not a finished manuscript. Bracketed text, planning language, "
         "and Word review comments must be resolved before submission. Equation objects are native "
         "Word mathematics, not screenshots; equation numbers use SEQ fields. Use Word headings "
         "for navigation and Review → New Comment for discussion.")
    body("Template provenance: adapted from the generic ACS Microsoft Word 2010 template "
         "downloaded from the ACS official server on 7 October 2026. It is not a JCTC-exclusive "
         "template. The original ACS paragraph styles are retained; review formatting is single "
         "column with 12-point Times New Roman text. The working draft is deliberately annotated; "
         "the initial-submission file must be clean.")
    body("JCTC author guidelines: "+GUIDELINES)
    body("Official ACS Word template: "+TEMPLATE_URL)
    body("Code mapping checked for this outline: ISCF=5 conventional XMVB-style L-BFGS; "
         "ISCF=7 TNHVP; ISCF=8 block-LBFGS; ISCF=9 NEO. Historical theory notes are not treated "
         "as proof of current implementation behavior. No new molecular calculations were performed.")
    body("Editorial decisions: (i) freeze the primary method and exact response policy; "
         "(ii) resolve endpoint validation for difficult OEO cases; (iii) freeze the benchmark "
         "version and matched inputs; (iv) select figures only after complete data are available; "
         "(v) agree authorship, funding, data deposition, and final novelty claims.")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    with zipfile.ZipFile(args.output) as package:
        document = package.read("word/document.xml")
        assert document.count(b"<m:oMath>") == equation_count
        assert b"SEQ Equation" in document
        assert "word/comments.xml" in package.namelist()
        assert b"wordprocessingml.document.main+xml" in package.read("[Content_Types].xml")
    print(f"{args.output}\nNative equations: {equation_count}; comments: {len(doc.comments)}")


if __name__ == "__main__":
    main()
